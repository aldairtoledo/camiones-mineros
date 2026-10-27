"""Reporte Excel profesional del seguimiento de camiones.

Arma un libro con estas hojas, en orden de lectura:
  Resumen      -> numeros clave, hallazgos y graficos (portada)
  Camiones     -> una fila por camion: tiempos, detenciones, distancia, velocidades, zonas
  Detenciones  -> cada parada de cada camion (inicio, fin, duracion, lugar)
  Zonas        -> uso de cada zona
  Ocupacion    -> camiones dentro de cada zona a lo largo del tiempo
  Bitacora     -> eventos y avisos en orden cronologico
  Posiciones   -> posicion de cada camion cada medio segundo
  Glosario     -> que significa cada termino y como se hizo el analisis

Solo usa openpyxl. Los graficos son nativos de Excel (se pueden editar).
Lo usan la app de escritorio y, mas adelante, la version web.
"""
import math
from collections import Counter, defaultdict
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as FuenteGrafico, Paragraph, \
    ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ENCABEZADO = "LABORATORIO DE SOFTWARE MINERO – UNI"
TITULO = "Reporte de seguimiento de camiones"
LETRA = "Arial"

# Colores (sobrios). Las series usan una paleta validada para daltonismo, en orden fijo.
C_OSCURO = "1E2430"
C_CABECERA = "2B3442"
C_ACENTO = "F2A900"
C_TEXTO = "1F2937"
C_SUAVE = "6B7280"
C_LINEA = "D9DDE3"
C_TARJETA = "F4F5F7"
C_FILA_ALTERNA = "F9FAFB"
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
C_MOVIMIENTO = SERIES[0]          # azul
C_DETENIDO = SERIES[1]            # naranja
C_DETENIDO_FUERA = SERIES[6]      # violeta
NIVEL_FONDO = {"Alerta": "FDE2E2", "Aviso": "FFF1CC", "Info": None}

# Una detencion es una parada de al menos este tiempo. Movimientos mas cortos que la
# tolerancia (un acomodo en la cola) no la cortan.
DETENCION_MINIMA_S = 5.0
TOLERANCIA_MOVIMIENTO_S = 2.0
# Con menos movimiento que esto, la velocidad media no es representativa y se deja vacia.
MOVIMIENTO_MINIMO_VEL_S = 3.0

TIPOS_AVISO = {
    "DETENIDO": "Detención", "ARRANQUE": "Arranque", "PASO_CERCANO": "Paso cercano",
    "PASO_CERCA_COLA": "Paso junto a la cola", "ENTRADA_ZONA": "Entrada a zona",
    "SALIDA_PANTALLA": "Salida de la imagen", "ENTRADA_PANTALLA": "Ingreso a la imagen",
    "ACERCAMIENTO": "Acercamiento a zona", "CURVA": "Curva", "REGRESO": "Regreso",
    "RETROCESO": "Retroceso", "GIRO_180": "Giro en U", "EXCESO_VELOCIDAD": "Exceso de velocidad",
}
NIVELES = {"info": "Info", "aviso": "Aviso", "alerta": "Alerta"}

_borde_fino = Side(style="thin", color=C_LINEA)
BORDE_CELDA = Border(bottom=_borde_fino)


# ------------------------------------------------------------------ utilidades
def _dias(segundos):
    """Segundos -> fraccion de dia (asi Excel lo trata como tiempo y se puede sumar)."""
    return None if segundos is None else segundos / 86400.0


def _formato_tiempo(duracion_s):
    return "[h]:mm:ss" if duracion_s >= 3600 else "[m]:ss"


def _hms(segundos):
    s = int(round(segundos or 0))
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60}:{s % 60:02d}"


def _f(tamano=10, negrita=False, color=C_TEXTO, cursiva=False):
    return Font(name=LETRA, size=tamano, bold=negrita, color=color, italic=cursiva)


def _relleno(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _texto_grafico(texto, tamano=11, negrita=True, color=C_TEXTO):
    cp = CharacterProperties(sz=tamano * 100, b=negrita, solidFill=color, latin=FuenteGrafico(typeface=LETRA))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=texto)])])


def _estilo_texto_ejes(tamano=9):
    cp = CharacterProperties(sz=tamano * 100, solidFill=C_SUAVE, latin=FuenteGrafico(typeface=LETRA))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(grafico, texto):
    grafico.title = Title(tx=Text(rich=_texto_grafico(texto)), overlay=False)


def _titulo_eje(eje, texto):
    eje.title = Title(tx=Text(rich=_texto_grafico(texto, 9, False, C_SUAVE)), overlay=False)


def _ejes_limpios(grafico, rejilla_y=True):
    """Ejes visibles y discretos: rejilla horizontal gris clara, sin rejilla vertical."""
    for eje in (grafico.x_axis, grafico.y_axis):
        eje.delete = False
        eje.txPr = _estilo_texto_ejes()
        eje.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_LINEA))
    grafico.x_axis.majorGridlines = None
    if rejilla_y:
        grafico.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E8EAED")))
    else:
        grafico.y_axis.majorGridlines = None
    if grafico.legend is not None:
        grafico.legend.position = "b"
        grafico.legend.txPr = _estilo_texto_ejes()


def _color_serie(serie, color, linea=False, ancho=22225):
    if linea:
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = ancho
    else:
        serie.graphicalProperties.solidFill = color
        serie.graphicalProperties.line.solidFill = "FFFFFF"     # separacion fina entre barras
        serie.graphicalProperties.line.width = 9525


# ------------------------------------------------------------------ hojas: comun
def _cabecera_hoja(ws, titulo, descripcion, ultima_col):
    """Titulo y descripcion arriba de cada hoja de datos (filas 1-2); la tabla empieza en la 4."""
    ws.sheet_view.showGridLines = False
    ws["A1"] = titulo
    ws["A1"].font = _f(14, True)
    ws["A2"] = descripcion
    ws["A2"].font = _f(10, False, C_SUAVE, True)
    ws.row_dimensions[1].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ultima_col, 6))
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"


def _tabla(ws, columnas, filas, fila_ini=4, alternar=True):
    """columnas: [(titulo, ancho, formato)]; filas: listas de valores. Devuelve la ultima fila."""
    for j, (titulo, ancho, _) in enumerate(columnas, 1):
        c = ws.cell(row=fila_ini, column=j, value=titulo)
        c.font = _f(10, True, "FFFFFF")
        c.fill = _relleno(C_CABECERA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila_ini].height = 48
    r = fila_ini
    for i, fila in enumerate(filas):
        r = fila_ini + 1 + i
        fondo = _relleno(C_FILA_ALTERNA) if alternar and i % 2 else None
        for j, valor in enumerate(fila, 1):
            c = ws.cell(row=r, column=j, value=valor)
            c.font = _f(10)
            c.border = BORDE_CELDA
            formato = columnas[j - 1][2]
            if formato:
                c.number_format = formato
            if fondo is not None:
                c.fill = fondo
            if isinstance(valor, str):
                c.alignment = Alignment(vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=fila_ini + 1, column=2)
    if filas:
        ws.auto_filter.ref = f"A{fila_ini}:{get_column_letter(len(columnas))}{r}"
    return r


def _sin_datos(ws, fila, texto="Sin datos registrados en este análisis."):
    ws.cell(row=fila, column=1, value=texto).font = _f(10, False, C_SUAVE, True)


# ------------------------------------------------------------------ calculos
class _Datos:
    """Ordena todo lo que llega del registro y calcula lo derivado (detenciones, etc.)."""

    def __init__(self, d):
        self.d = d
        self.video = d.get("video") or "-"
        self.duracion = float(d.get("duracion_s") or 0.0)
        self.ftiempo = _formato_tiempo(self.duracion)
        self.georef = d.get("georef")
        self.metros = self.georef is not None
        self.intervalo = float(d.get("intervalo") or 0.5)
        self.nombres = list(d.get("nombres_zonas") or [])
        self.tipos = list(d.get("tipos_zonas") or [])
        self.camiones = list(d.get("camiones") or [])
        self.zonas = list(d.get("zonas") or [])
        self.eventos = list(d.get("eventos") or [])
        self.alertas = list(d.get("alertas") or [])
        orden = lambda f: (f["tiempo_s"], f["id_camion"])
        # Para calcular se usan todas las posiciones; la hoja puede ir vacia si son demasiadas.
        self.tray = sorted(d.get("trayectorias_calculo") or d.get("trayectorias") or [], key=orden)
        self.tray_hoja = self.tray if d.get("trayectorias") else []
        self.ocupacion = list(d.get("ocupacion") or [])
        self.limite_general = d.get("limite_general")
        self.unidad_d = "m" if self.metros else "px"
        self.unidad_v = "km/h" if self.metros else "px/s"
        self._posiciones()
        # Entradas y salidas por el borde de la imagen (de los avisos)
        self.salidas, self.ingresos = {}, {}
        for a in self.alertas:
            if a.get("tipo") == "SALIDA_PANTALLA":
                self.salidas.setdefault(a.get("id_camion"), a.get("tiempo_s"))
            elif a.get("tipo") == "ENTRADA_PANTALLA":
                self.ingresos.setdefault(a.get("id_camion"), a.get("tiempo_s"))
        self.detenciones = self._calcular_detenciones()
        self.por_camion = defaultdict(list)
        for det in self.detenciones:
            self.por_camion[det["camion"]].append(det)

    # -- zonas
    def nombre_zona(self, z):
        nombre = self.nombres[z].strip() if z < len(self.nombres) and self.nombres[z] else ""
        return nombre or f"Zona {z + 1}"

    def nombre_codigos(self, codigos):
        """'Z1, Z3' -> 'Chancadora, Zona 3'; '-' -> 'Fuera de zonas'."""
        if not codigos or codigos == "-":
            return "Fuera de zonas"
        nombres = []
        for c in str(codigos).split(","):
            c = c.strip()
            if c.startswith("Z") and c[1:].isdigit():
                nombres.append(self.nombre_zona(int(c[1:]) - 1))
        return ", ".join(nombres) or "Fuera de zonas"

    # -- posiciones en la unidad del reporte
    def _posiciones(self):
        """Columna X/Y que se muestra: metros del plano local, Este/Norte o pixeles."""
        if not self.tray:
            self.eje_x, self.eje_y, self.etiquetas_xy = "centro_x", "centro_y", ("Píxel X", "Píxel Y")
        elif "x_m" in self.tray[0]:
            self.eje_x, self.eje_y, self.etiquetas_xy = "x_m", "y_m", ("X (m)", "Y (m)")
        elif "este_m" in self.tray[0]:
            self.eje_x, self.eje_y, self.etiquetas_xy = "este_m", "norte_m", ("Este (m)", "Norte (m)")
        else:
            self.eje_x, self.eje_y, self.etiquetas_xy = "centro_x", "centro_y", ("Píxel X", "Píxel Y")

    # -- detenciones
    def _calcular_detenciones(self):
        por = defaultdict(list)
        for f in self.tray:
            por[f["id_camion"]].append(f)
        salida = []
        hueco_max = TOLERANCIA_MOVIMIENTO_S + self.intervalo + 1e-6
        for pid, filas in sorted(por.items()):
            ultima_t = filas[-1]["tiempo_s"]
            grupo = []

            def cerrar():
                if not grupo:
                    return
                ini, fin = grupo[0]["tiempo_s"], grupo[-1]["tiempo_s"] + self.intervalo
                dur = fin - ini
                if dur < DETENCION_MINIMA_S:
                    return
                lugar = Counter(g.get("zona_actual", "-") for g in grupo).most_common(1)[0][0]
                xs = [g[self.eje_x] for g in grupo if g.get(self.eje_x) is not None]
                ys = [g[self.eje_y] for g in grupo if g.get(self.eje_y) is not None]
                notas = []
                if grupo[0] is filas[0]:
                    notas.append("ya estaba detenido al aparecer")
                if grupo[-1]["tiempo_s"] >= ultima_t - 1e-6 and self.duracion - ultima_t <= 2 * self.intervalo + 0.5:
                    notas.append("seguía detenido al terminar el análisis")
                salida.append({"camion": pid, "inicio": ini, "fin": min(fin, self.duracion or fin),
                               "duracion": dur, "lugar": self.nombre_codigos(lugar),
                               "en_zona": lugar not in (None, "", "-"),
                               "x": sum(xs) / len(xs) if xs else None, "y": sum(ys) / len(ys) if ys else None,
                               "nota": "; ".join(notas)})

            for f in filas:
                if f.get("movimiento") == "DETENIDO":
                    if grupo and f["tiempo_s"] - grupo[-1]["tiempo_s"] > hueco_max:
                        cerrar()
                        grupo = []
                    grupo.append(f)
            cerrar()
        return sorted(salida, key=lambda x: (x["inicio"], x["camion"]))

    # -- reparto del tiempo de la flota (para la dona)
    def reparto_tiempo(self):
        mov = det_zona = det_fuera = 0
        for f in self.tray:
            if f.get("movimiento") == "EN_MOVIMIENTO":
                mov += 1
            elif f.get("zona_actual", "-") not in ("-", "", None):
                det_zona += 1
            else:
                det_fuera += 1
        return [("En movimiento", mov * self.intervalo), ("Detenido en una zona", det_zona * self.intervalo),
                ("Detenido fuera de zonas", det_fuera * self.intervalo)]

    # -- ocupacion por intervalos
    def ocupacion_intervalos(self):
        if self.duracion <= 600:
            paso = 10
        elif self.duracion <= 1800:
            paso = 30
        else:
            paso = 60
        acum = defaultdict(list)
        zonas = sorted({o["zona"] for o in self.ocupacion}, key=lambda z: int(z[1:]))
        for o in self.ocupacion:
            acum[(int(o["tiempo_s"] // paso), o["zona"])].append(o["camiones"])
        bloques = sorted({b for b, _ in acum})
        filas = []
        for b in bloques:
            fila = [b * paso]
            for z in zonas:
                v = acum.get((b, z), [])
                fila += [sum(v) / len(v) if v else 0.0, max(v) if v else 0]
            filas.append(fila)
        return paso, zonas, filas


# ------------------------------------------------------------------ hoja Resumen
def _kpis(datos, ref):
    """Tarjetas: (titulo, formula o valor, formato, nota)."""
    n = len(datos.camiones)
    tot = ref["fila_total"]
    col = ref["col"]
    cam = "'Camiones'!"
    t = []
    t.append(("Camiones analizados", n, "0", f"en {_hms(datos.duracion)} de video"))
    if "pct_det" not in col:              # sin camiones: tarjetas en cero
        return t + [("Detenciones", 0, "0", "sin datos"), ("Zonas de control", len(datos.zonas), "0",
                                                           "definidas en el video")]
    if col.get("distancia"):
        t.append((f"Distancia total recorrida", f"={cam}{col['distancia']}{tot}",
                  '#,##0 "m"' if datos.metros else '#,##0 "px"', "suma de todos los camiones"))
    if col.get("vel_media"):
        t.append(("Velocidad media en movimiento", f"={cam}{col['vel_media']}{tot}",
                  f'0.0 "{datos.unidad_v}"', "sin contar el tiempo detenido"))
    t.append(("Tiempo detenido de la flota", f"={cam}{col['pct_det']}{tot}", "0%",
              "del tiempo total observado"))
    t.append(("Detenciones", f"={cam}{col['detenciones']}{tot}", "0",
              f"paradas de {DETENCION_MINIMA_S:g} s o más"))
    if col.get("excesos"):
        nota = f"límite general {datos.limite_general:g} km/h" if datos.limite_general else "según los límites definidos"
        t.append(("Excesos de velocidad", f"={cam}{col['excesos']}{tot}", "0", nota))
    else:
        t.append(("Zonas de control", len(datos.zonas), "0", "definidas en el video"))
    return t[:6]


def _hallazgos(datos):
    """Frases cortas, en lenguaje simple, con lo mas importante del analisis."""
    h = []
    cam = [c for c in datos.camiones if c.get("tiempo_visible_s")]
    if cam:
        casi_quietos = [c for c in cam if c.get("porcentaje_detenido", 0) >= 90]
        if len(casi_quietos) >= 2:
            lugares = Counter()
            for c in casi_quietos:
                for det in datos.por_camion.get(c["id_camion"], []):
                    lugares[det["lugar"]] += det["duracion"]
            donde = f", principalmente en {lugares.most_common(1)[0][0]}" if lugares else ""
            h.append(f"{len(casi_quietos)} de {len(cam)} camiones estuvieron detenidos más del 90 % del tiempo"
                     f"{donde} (camiones: {', '.join(str(c['id_camion']) for c in casi_quietos)}).")
        else:
            peor = max(cam, key=lambda c: c.get("porcentaje_detenido", 0))
            lugares = Counter()
            for det in datos.por_camion.get(peor["id_camion"], []):
                lugares[det["lugar"]] += det["duracion"]
            donde = f", sobre todo en {lugares.most_common(1)[0][0]}" if lugares else ""
            h.append(f"El camión {peor['id_camion']} fue el que más tiempo pasó detenido: "
                     f"{peor.get('porcentaje_detenido', 0):.0f} % del tiempo que se vio "
                     f"({_hms(peor.get('tiempo_detenido_s'))}){donde}.")
        moviles = [c for c in cam if c.get("porcentaje_detenido", 100) < 90]
        if len(casi_quietos) >= 2 and moviles:
            h.append(f"Los camiones que sí trabajaron ({', '.join(str(c['id_camion']) for c in moviles)}) estuvieron "
                     f"en movimiento entre {min(100 - c['porcentaje_detenido'] for c in moviles):.0f} % y "
                     f"{max(100 - c['porcentaje_detenido'] for c in moviles):.0f} % del tiempo.")
    if datos.detenciones:
        larga = max(datos.detenciones, key=lambda x: x["duracion"])
        h.append(f"La detención más larga fue del camión {larga['camion']}: {_hms(larga['duracion'])} "
                 f"en {larga['lugar']} (empezó en el {_hms(larga['inicio'])} del video).")
    zonas = [z for z in datos.zonas if z.get("ocupacion_maxima")]
    if zonas:
        z = max(zonas, key=lambda z: (z["ocupacion_maxima"], z.get("ocupacion_promedio") or 0))
        n = z["ocupacion_maxima"]
        h.append(f"La zona con más camiones a la vez fue {z['nombre_zona']}: hasta {n} "
                 f"{'camión' if n == 1 else 'camiones'} (en promedio {z.get('ocupacion_promedio') or 0:.1f}).")
    if datos.metros and cam:
        lejos = max(cam, key=lambda c: c.get("distancia_m", 0) or 0)
        if (lejos.get("distancia_m") or 0) > 0:
            metros = f"{lejos['distancia_m']:,.0f}".replace(",", " ")      # 1 234 (solo el numero)
            h.append(f"El camión {lejos['id_camion']} fue el que más recorrió: {metros} m, "
                     f"a {lejos.get('velocidad_media_en_movimiento_km_h', 0):.0f} km/h en promedio "
                     f"mientras avanzaba.")
    n_excesos = sum(c.get("veces_sobre_el_limite") or 0 for c in cam) if any(
        "veces_sobre_el_limite" in c for c in cam) else sum(1 for a in datos.alertas if a.get("tipo") == "EXCESO_VELOCIDAD")
    if n_excesos:
        rapido = max(cam, key=lambda c: c.get("velocidad_maxima_km_h", 0) or 0) if cam else None
        extra = (f"; la velocidad más alta fue del camión {rapido['id_camion']} "
                 f"({rapido['velocidad_maxima_km_h']:.0f} km/h)") if rapido and rapido.get("velocidad_maxima_km_h") else ""
        h.append(f"Se registraron {n_excesos} excesos de velocidad{extra}.")
    elif datos.metros and datos.limite_general:
        h.append(f"Ningún camión superó el límite de {datos.limite_general:g} km/h.")
    if datos.salidas:
        al_inicio = {c["id_camion"] for c in cam if (c.get("primera_aparicion_s") or 0) <= 2.0}
        nuevos = sorted(p for p in datos.ingresos if p not in al_inicio)
        n = len(datos.salidas)
        texto = (f"{n} {'camión salió' if n == 1 else 'camiones salieron'} de la imagen "
                 f"({', '.join(str(p) for p in sorted(datos.salidas))})")
        if nuevos:
            texto += (f"; los que entraron por el borde recibieron un número nuevo "
                      f"({', '.join(str(p) for p in nuevos)})")
        h.append(texto + ".")
    estimadas = [c.get("porcentaje_posiciones_estimadas", 0) for c in cam]
    if estimadas and sum(estimadas) / len(estimadas) >= 30:
        h.append(f"El {sum(estimadas) / len(estimadas):.0f} % de las posiciones fueron estimadas (camiones juntos "
                 "o tapados). Los tiempos son confiables; las distancias de esos tramos, aproximadas.")
    if not h:
        h.append("No hubo suficientes datos para sacar conclusiones.")
    return h[:8]


def _hoja_resumen(wb, datos, ref, graficos):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 13.5
    ws.column_dimensions["N"].width = 2
    # --- banda superior
    for r in range(1, 5):
        for j in range(1, 15):
            ws.cell(row=r, column=j).fill = _relleno(C_OSCURO)
    ws["B1"] = ENCABEZADO
    ws["B1"].font = _f(9, True, C_ACENTO)
    ws["B2"] = TITULO
    ws["B2"].font = _f(20, True, "FFFFFF")
    calib = "sin calibrar (medidas en píxeles)"
    if datos.georef is not None:
        info = getattr(datos.georef, "info", {}) or {}
        pct = info.get("incertidumbre_escala_pct")
        calib = "en metros" + (f" (± {pct:g} %)" if pct else "")
        if str(info.get("fuente", "")).startswith(("Automática", "Solo camión")):
            calib += ", aproximada con el tamaño del camión"
        if info.get("revisar"):
            calib += " — revisar calibración"
    ws["B3"] = (f"Video: {datos.video}   ·   Analizado: {_hms(datos.duracion)} de video   ·   "
                f"Calibración: {calib}   ·   Generado: {datetime.now():%d/%m/%Y %H:%M}")
    ws["B3"].font = _f(9, False, "C9CED6")
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 8

    # --- tarjetas
    fila = 6
    for i, (titulo, valor, formato, nota) in enumerate(_kpis(datos, ref)):
        c0 = 2 + i * 2
        for r in (fila, fila + 1, fila + 2):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
            for j in (c0, c0 + 1):
                celda = ws.cell(row=r, column=j)
                celda.fill = _relleno(C_TARJETA)
                celda.border = Border(left=Side(style="medium", color="FFFFFF") if j == c0 else None,
                                      right=Side(style="medium", color="FFFFFF") if j == c0 + 1 else None,
                                      top=Side(style="thick", color=C_ACENTO) if r == fila else None)
        a = ws.cell(row=fila, column=c0, value=titulo.upper())
        a.font = _f(8, True, C_SUAVE)
        a.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        v = ws.cell(row=fila + 1, column=c0, value=valor)
        v.font = _f(20, True)
        v.number_format = formato
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        nn = ws.cell(row=fila + 2, column=c0, value=nota)
        nn.font = _f(8, False, C_SUAVE)
        nn.alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)
    ws.row_dimensions[fila].height = 22
    ws.row_dimensions[fila + 1].height = 32
    ws.row_dimensions[fila + 2].height = 26

    # --- hallazgos y guia de lectura
    fila = 10
    ws.cell(row=fila, column=2, value="Hallazgos principales").font = _f(12, True)
    ws.cell(row=fila, column=9, value="Cómo leer este reporte").font = _f(12, True)
    for j in range(2, 14):
        ws.cell(row=fila, column=j).border = Border(bottom=Side(style="medium", color=C_ACENTO))
    guia = [("Resumen", "números clave y gráficos (esta hoja)"),
            ("Camiones", "una fila por camión: tiempos, distancia, velocidad"),
            ("Detenciones", f"cada parada de {DETENCION_MINIMA_S:g} s o más, con hora y lugar"),
            ("Zonas", "entradas y permanencia en cada zona"),
            ("Ocupación", "cuántos camiones había en cada zona; cuándo hay cola"),
            ("Bitácora", "todo lo que pasó, en orden de tiempo"),
            ("Posiciones", f"dónde estaba cada camión cada {datos.intervalo:g} s"),
            ("Glosario", "qué significa cada término")]
    hallazgos = _hallazgos(datos)
    for i in range(max(len(hallazgos), len(guia))):
        r = fila + 1 + i
        ws.row_dimensions[r].height = 30
        if i < len(hallazgos):
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            c = ws.cell(row=r, column=2, value=f"▸  {hallazgos[i]}")
            c.font = _f(10)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if i < len(guia):
            ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
            h = ws.cell(row=r, column=9, value=guia[i][0])
            h.font = _f(10, True)
            h.alignment = Alignment(vertical="center")
            g = ws.cell(row=r, column=10, value=guia[i][1])
            g.font = _f(10, False, C_SUAVE)
            g.alignment = Alignment(vertical="center", wrap_text=True)
    fila = fila + 2 + max(len(hallazgos), len(guia))

    # --- graficos (2 por fila)
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=fila - 1))
    ws.cell(row=fila, column=2, value="Gráficos").font = _f(12, True)
    for j in range(2, 14):
        ws.cell(row=fila, column=j).border = Border(bottom=Side(style="medium", color=C_ACENTO))
    fila += 2
    for k, grafico in enumerate(graficos):
        col = "B" if k % 2 == 0 else "H"
        grafico.width, grafico.height = 15.6, 8.2
        ws.add_chart(grafico, f"{col}{fila}")
        if k % 2 == 1:
            fila += 18
            if k % 4 == 3:
                ws.row_breaks.append(Break(id=fila - 1))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


# ------------------------------------------------------------------ hoja Camiones
def _hoja_camiones(wb, datos):
    ws = wb.create_sheet("Camiones")
    ft = datos.ftiempo
    cols = [("Camión", 9, "0"), ("Primera aparición", 11, ft), ("Última aparición", 11, ft),
            ("Cómo apareció", 15, None), ("Salió de la imagen", 11, ft),
            ("Tiempo visible", 11, ft), ("En movimiento", 11, ft), ("Detenido", 11, ft),
            ("% detenido", 10, "0%"), ("N.º de detenciones", 11, "0"), ("Detención más larga", 11, ft)]
    if datos.metros:
        cols += [("Distancia recorrida (m)", 12, "#,##0"), ("Velocidad media en movimiento (km/h)", 13, "0.0")]
        hay_vel = any("velocidad_maxima_km_h" in c for c in datos.camiones)
        if hay_vel:
            cols += [("Velocidad máxima (km/h)", 11, "0.0"), ("Excesos de velocidad", 10, "0"),
                     ("Tiempo sobre el límite", 11, ft)]
    else:
        cols += [("Distancia recorrida (px)", 12, "#,##0"), ("Velocidad media en movimiento (px/s)", 13, "0.0")]
        hay_vel = False
    nz = len(datos.zonas)
    zona_cols = []
    for z in range(nz):
        nombre = datos.nombre_zona(z)
        hay_ciclo = any(c.get(f"ciclo_prom_Z{z + 1}_s") for c in datos.camiones)
        zona_cols.append((z, hay_ciclo))
        cols += [(f"Entradas a {nombre}", 10, "0"), (f"Tiempo en {nombre}", 11, ft)]
        if hay_ciclo:
            cols.append((f"Ciclo promedio en {nombre}", 11, ft))
    cols.append(("Posiciones estimadas", 11, "0%"))

    filas = []
    for c in datos.camiones:
        pid = c["id_camion"]
        dets = datos.por_camion.get(pid, [])
        primera = c.get("primera_aparicion_s") or 0
        como = ("Al inicio del video" if primera <= 2.0 else
                "Entró por el borde" if pid in datos.ingresos else "Apareció en la imagen")
        salio = _dias(datos.salidas[pid]) if pid in datos.salidas else "No"
        fila = [pid, _dias(c.get("primera_aparicion_s")), _dias(c.get("ultima_aparicion_s")), como, salio,
                _dias(c.get("tiempo_visible_s")), _dias(c.get("tiempo_en_movimiento_s")),
                _dias(c.get("tiempo_detenido_s")),
                (c.get("tiempo_detenido_s") or 0) / c["tiempo_visible_s"] if c.get("tiempo_visible_s") else 0,
                len(dets), _dias(max((x["duracion"] for x in dets), default=0))]
        poco = (c.get("tiempo_en_movimiento_s") or 0) < MOVIMIENTO_MINIMO_VEL_S
        if datos.metros:
            fila += [c.get("distancia_m", 0), None if poco else c.get("velocidad_media_en_movimiento_km_h", 0)]
            if hay_vel:
                fila += [c.get("velocidad_maxima_km_h"), c.get("veces_sobre_el_limite"),
                         _dias(c.get("tiempo_sobre_el_limite_s"))]
        else:
            fila += [c.get("distancia_px", 0), None if poco else c.get("velocidad_media_en_movimiento_px_s", 0)]
        for z, hay_ciclo in zona_cols:
            fila += [c.get(f"entradas_Z{z + 1}", 0), _dias(c.get(f"tiempo_total_en_Z{z + 1}_s", 0))]
            if hay_ciclo:
                fila.append(_dias(c.get(f"ciclo_prom_Z{z + 1}_s")))
        fila.append((c.get("porcentaje_posiciones_estimadas") or 0) / 100)
        filas.append(fila)

    _cabecera_hoja(ws, "Camiones",
                   "Una fila por camión. Los tiempos están en minutos:segundos. La última fila resume toda la flota. "
                   "Distancia y velocidad media solo cuentan los tramos en que el camión avanzaba. Un camión que "
                   "entra por el borde siempre recibe un número nuevo.", len(cols))
    ultima = _tabla(ws, cols, filas)
    ref = {"col": {}, "fila_total": ultima + 1}
    L = get_column_letter
    # Columna de cada dato, buscada por su titulo (asi no importa el orden).
    pos = {titulo: k for k, (titulo, _, _) in enumerate(cols, 1)}
    def col(prefijo):
        return next((k for t, k in pos.items() if t.startswith(prefijo)), None)
    c_vis, c_mov, c_det = col("Tiempo visible"), col("En movimiento"), col("Detenido")
    c_pct, c_ndet, c_larga = col("% detenido"), col("N.º de detenciones"), col("Detención más larga")
    c_dist, c_vel = col("Distancia recorrida"), col("Velocidad media")
    c_vmax, c_exc, c_texc = col("Velocidad máxima"), col("Excesos"), col("Tiempo sobre el límite")
    if filas:
        tot = ultima + 1
        r0, r1 = 5, ultima
        ws.cell(row=tot, column=1, value="Flota")
        for k in range(1, len(cols) + 1):
            c = ws.cell(row=tot, column=k)
            c.font = _f(10, True)
            c.fill = _relleno("E9ECF1")
            c.border = Border(top=Side(style="medium", color=C_CABECERA))
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = cols[k - 1][2] or "General"
            titulo = cols[k - 1][0]
            rango = f"{L(k)}{r0}:{L(k)}{r1}"
            if titulo.startswith("Primera"):
                c.value = f"=MIN({rango})"
            elif titulo.startswith("Última") or k in (c_larga, c_vmax):
                c.value = f"=MAX({rango})"
            elif titulo.startswith("Salió"):
                c.value = f'=IF(COUNT({rango})=1,"1 salió",COUNT({rango})&" salieron")'
                c.number_format = "General"
            elif k == c_pct:
                c.value = f"=IF({L(c_vis)}{tot}>0,{L(c_det)}{tot}/{L(c_vis)}{tot},0)"
            elif k == c_vel:
                # velocidad media de la flota = distancia total / tiempo total en movimiento
                factor = "*3.6" if datos.metros else ""
                c.value = f"=IF({L(c_mov)}{tot}>0,{L(c_dist)}{tot}/({L(c_mov)}{tot}*86400){factor},0)"
            elif titulo.startswith("Ciclo") or titulo.startswith("Posiciones"):
                c.value = f"=AVERAGE({rango})"
            elif k in (c_vis, c_mov, c_det, c_ndet, c_dist, c_exc, c_texc) or titulo.startswith(("Entradas", "Tiempo en")):
                c.value = f"=SUM({rango})"
        ws.cell(row=tot, column=1).alignment = Alignment(horizontal="left", vertical="center")
        # Barras dentro de la celda (% detenido) y excesos en rojo
        ws.conditional_formatting.add(f"{L(c_pct)}{r0}:{L(c_pct)}{r1}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=C_DETENIDO))
        if c_exc:
            ws.conditional_formatting.add(f"{L(c_exc)}{r0}:{L(c_exc)}{r1}", CellIsRule(
                operator="greaterThan", formula=["0"], font=Font(name=LETRA, bold=True, color="C53030"),
                fill=_relleno("FDE2E2")))
        ref["col"] = {"distancia": L(c_dist), "vel_media": L(c_vel), "pct_det": L(c_pct), "detenciones": L(c_ndet)}
        if c_exc:
            ref["col"]["excesos"] = L(c_exc)
    else:
        _sin_datos(ws, 5)
    ws.page_setup.orientation = "landscape"
    return ref


# ------------------------------------------------------------------ hoja Detenciones
def _hoja_detenciones(wb, datos):
    ws = wb.create_sheet("Detenciones")
    ft = datos.ftiempo
    ex, ey = datos.etiquetas_xy
    formato_xy = "0.0" if datos.metros else "0"
    cols = [("N.º", 6, "0"), ("Camión", 9, "0"), ("Inicio", 10, ft), ("Fin", 10, ft), ("Duración", 10, ft),
            ("Lugar", 24, None), (f"Posición {ex}", 11, formato_xy), (f"Posición {ey}", 11, formato_xy),
            ("Observación", 48, None)]
    filas = [[i, d["camion"], _dias(d["inicio"]), _dias(d["fin"]), _dias(d["duracion"]), d["lugar"],
              d["x"], d["y"], d["nota"] or ""] for i, d in enumerate(datos.detenciones, 1)]
    _cabecera_hoja(ws, "Detenciones",
                   f"Cada vez que un camión estuvo parado {DETENCION_MINIMA_S:g} s o más. Movimientos de menos de "
                   f"{TOLERANCIA_MOVIMIENTO_S:g} s (por ejemplo, acomodarse en la cola) no cortan la detención. "
                   "La posición es el punto promedio donde estuvo parado.", len(cols))
    ultima = _tabla(ws, cols, filas)
    if filas:
        ws.conditional_formatting.add(f"E5:E{ultima}", DataBarRule(start_type="num", start_value=0,
                                                                    end_type="max", color=C_DETENIDO))
    else:
        _sin_datos(ws, 5, f"No hubo detenciones de {DETENCION_MINIMA_S:g} s o más.")


# ------------------------------------------------------------------ hoja Zonas
def _hoja_zonas(wb, datos):
    ws = wb.create_sheet("Zonas")
    ft = datos.ftiempo
    cols = [("Código", 8, None), ("Zona", 20, None), ("Tipo", 14, None), ("Qué mide", 30, None),
            ("Entradas", 10, "0"), ("Camiones por hora", 11, "0.0"), ("Permanencia promedio", 12, ft),
            ("Permanencia mínima", 12, ft), ("Permanencia máxima", 12, ft),
            ("Ocupación promedio (camiones)", 12, "0.0"), ("Ocupación máxima (camiones)", 12, "0")]
    filas = []
    for z in datos.zonas:
        tipo = str(z.get("tipo", "")).lower()
        filas.append([z.get("zona"), z.get("nombre_zona"),
                      "Acumulativa" if tipo == "acumulativa" else "Dinámica",
                      "Cuenta los camiones que pasan" if tipo == "acumulativa" else "Cuenta los camiones que hay dentro",
                      z.get("entradas_totales"), z.get("camiones_por_hora"),
                      _dias(z.get("permanencia_prom_s")), _dias(z.get("permanencia_min_s")),
                      _dias(z.get("permanencia_max_s")), z.get("ocupacion_promedio"), z.get("ocupacion_maxima")])
    _cabecera_hoja(ws, "Zonas",
                   "Uso de cada zona de control. Permanencia = cuánto tiempo estuvo un camión dentro en cada visita. "
                   "Ocupación = cuántos camiones había dentro al mismo tiempo. Camiones por hora se calcula con el "
                   "tiempo analizado.", len(cols))
    _tabla(ws, cols, filas)
    if not filas:
        _sin_datos(ws, 5, "No se definieron zonas.")


# ------------------------------------------------------------------ hoja Ocupacion
def _hoja_ocupacion(wb, datos):
    ws = wb.create_sheet("Ocupación")
    paso, zonas, filas = datos.ocupacion_intervalos()
    ft = datos.ftiempo
    cols = [("Desde", 10, ft)]
    for z in zonas:
        nombre = datos.nombre_zona(int(z[1:]) - 1)
        cols += [(f"{nombre}: promedio", 13, "0.0"), (f"{nombre}: máximo", 12, "0")]
    _cabecera_hoja(ws, "Ocupación de zonas en el tiempo",
                   f"Camiones dentro de cada zona, en intervalos de {paso} s. Sirve para ver cuándo se forma cola "
                   "(por ejemplo, frente a la chancadora).", len(cols))
    _tabla(ws, cols, [[_dias(f[0])] + f[1:] for f in filas])
    if not filas:
        _sin_datos(ws, 5)
    return paso, zonas, filas


# ------------------------------------------------------------------ hoja Bitacora
def _hoja_bitacora(wb, datos):
    ws = wb.create_sheet("Bitácora")
    filas = []
    for e in datos.eventos:
        ev = e.get("evento")
        zona = e.get("nombre_zona") or e.get("zona")
        if ev == "ENTRADA":
            tipo, desc = "Entrada a zona", f"El camión {e['id_camion']} entró a {zona}."
        elif ev == "SALIDA":
            perm = e.get("permanencia_s")
            tipo = "Salida de zona"
            desc = f"El camión {e['id_camion']} salió de {zona}" + (f" (estuvo {_hms(perm)})." if perm else ".")
        else:
            tipo, desc = "Sigue en la zona", (f"Al terminar, el camión {e['id_camion']} seguía en {zona} "
                                              f"(lleva {_hms(e.get('permanencia_s'))}).")
        filas.append((e["tiempo_s"], e["id_camion"], tipo, zona, desc, "Info"))
    for a in datos.alertas:
        if a.get("tipo") == "ENTRADA_ZONA":       # ya esta en los eventos de zona
            continue
        filas.append((a.get("tiempo_s", 0), a.get("id_camion"), TIPOS_AVISO.get(a.get("tipo"), a.get("titulo", "")),
                      a.get("zona") or "", a.get("detalle", ""), NIVELES.get(a.get("nivel"), "Info")))
    filas.sort(key=lambda f: (f[0], f[1] or 0))
    cols = [("Hora en el video", 11, datos.ftiempo), ("Camión", 9, "0"), ("Evento", 22, None),
            ("Zona", 18, None), ("Descripción", 62, None), ("Importancia", 12, None)]
    _cabecera_hoja(ws, "Bitácora",
                   "Todo lo que pasó durante el video, en orden de tiempo: entradas y salidas de zonas, "
                   "detenciones, arranques, excesos de velocidad y otros avisos.", len(cols))
    ultima = _tabla(ws, cols, [[_dias(t), pid, tipo, zona, desc, nivel] for t, pid, tipo, zona, desc, nivel in filas])
    for r in range(5, ultima + 1):
        nivel = ws.cell(row=r, column=6).value
        fondo = NIVEL_FONDO.get(nivel)
        if fondo:
            ws.cell(row=r, column=6).fill = _relleno(fondo)
            ws.cell(row=r, column=6).font = _f(10, True, "9B2C2C" if nivel == "Alerta" else "8A5A00")
    if not filas:
        _sin_datos(ws, 5)


# ------------------------------------------------------------------ hoja Posiciones
def _hoja_posiciones(wb, datos, aviso=None):
    ws = wb.create_sheet("Posiciones")
    ex, ey = datos.etiquetas_xy
    fxy = "0.00" if datos.metros else "0.0"
    cols = [("Tiempo (s)", 9, "0.00"), ("Hora en el video", 10, datos.ftiempo), ("Camión", 8, "0")]
    con_metros = datos.eje_x != "centro_x"
    if con_metros:
        cols += [(ex, 10, fxy), (ey, 10, fxy)]
    extra_en = datos.tray and "este_m" in datos.tray[0] and datos.eje_x == "x_m"
    if extra_en:
        cols += [("Este (m)", 12, "0.00"), ("Norte (m)", 12, "0.00")]
    if datos.metros:
        cols.append(("Velocidad (km/h)", 10, "0.0"))
    cols += [("Marcha", 13, None), ("Zona", 18, None), ("Posición", 11, None),
             ("Píxel X", 9, "0"), ("Píxel Y", 9, "0")]
    filas = []
    for f in datos.tray_hoja:
        fila = [f["tiempo_s"], _dias(f["tiempo_s"]), f["id_camion"]]
        if con_metros:
            fila += [f.get(datos.eje_x), f.get(datos.eje_y)]
        if extra_en:
            fila += [f.get("este_m"), f.get("norte_m")]
        if datos.metros:
            fila.append(f.get("velocidad_km_h"))
        fila += ["En movimiento" if f.get("movimiento") == "EN_MOVIMIENTO" else "Detenido",
                 datos.nombre_codigos(f.get("zona_actual")),
                 "Medida" if f.get("estado") == "DETECTADO" else "Estimada",
                 f.get("centro_x"), f.get("centro_y")]
        filas.append(fila)
    origen = ("X hacia la derecha e Y hacia arriba, desde la esquina inferior izquierda del video"
              if datos.eje_x == "x_m" else "píxeles de la imagen original (origen arriba a la izquierda)")
    _cabecera_hoja(ws, "Posiciones",
                   f"Dónde estaba cada camión cada {datos.intervalo:g} s ({origen}). 'Estimada' = el camión estaba "
                   "tapado o junto a otro y su posición se calculó. Usa los filtros del encabezado para ver un "
                   "solo camión.", len(cols))
    if aviso:
        ws.cell(row=3, column=1, value=aviso).font = _f(10, True, "9B2C2C")
    _tabla(ws, cols, filas, alternar=False)
    if not filas:
        _sin_datos(ws, 5)


# ------------------------------------------------------------------ hoja Glosario
def _hoja_glosario(wb, datos, parametros):
    ws = wb.create_sheet("Glosario")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    ws["A1"] = "Glosario y parámetros"
    ws["A1"].font = _f(14, True)
    ws["A2"] = "Qué significa cada término del reporte y con qué configuración se hizo el análisis."
    ws["A2"].font = _f(10, False, C_SUAVE, True)
    terminos = [
        ("Camión (N.º)", "Número que la app le asigna a cada camión al detectarlo. Se mantiene mientras el camión "
                         "esté en la imagen. Si un camión sale por el borde y luego entra uno, el que entra recibe "
                         "un número nuevo (un camión real tarda mucho más en completar su ciclo y volver)."),
        ("Tiempo visible", "Tiempo total en que el camión estuvo en la imagen."),
        ("En movimiento / Detenido", "Se considera detenido cuando prácticamente no avanza (solo el temblor normal "
                                     "de la detección)."),
        ("Detención", f"Parada de {DETENCION_MINIMA_S:g} s o más. Movimientos menores de "
                      f"{TOLERANCIA_MOVIMIENTO_S:g} s no la interrumpen."),
        ("Zona acumulativa", "Zona que cuenta cuántos camiones pasaron por ella (por ejemplo, la descarga)."),
        ("Zona dinámica", "Zona que cuenta cuántos camiones hay dentro en cada momento (por ejemplo, una cola)."),
        ("Permanencia", "Tiempo que un camión estuvo dentro de una zona en una visita."),
        ("Ocupación", "Cantidad de camiones dentro de una zona al mismo tiempo."),
        ("Camiones por hora", "Entradas a una zona, llevadas a una hora: entradas ÷ horas analizadas."),
        ("Ciclo", "Tiempo entre dos entradas seguidas del mismo camión a la misma zona (una vuelta completa)."),
        ("Velocidad media en movimiento", "Distancia recorrida ÷ tiempo en movimiento. No cuenta el tiempo detenido."),
        ("Exceso de velocidad", "El camión superó el límite (general o de la zona) durante al menos 1 s."),
        ("Posición estimada", "El camión estaba tapado o muy junto a otro; su posición se calculó a partir de su "
                              "movimiento anterior."),
        ("Calibración", "Conversión de píxeles a metros. Con distancias conocidas es precisa (± 1 %); con el "
                        "tamaño del camión es aproximada (± 5–10 %)."),
        ("Formato de tiempo", "minutos:segundos (por ejemplo, 2:05 = 2 minutos y 5 segundos)."),
    ]
    r = 4
    for titulo in ("Término", "Significado"):
        pass
    for j, titulo in enumerate(("Término", "Significado"), 1):
        c = ws.cell(row=r, column=j, value=titulo)
        c.font = _f(10, True, "FFFFFF")
        c.fill = _relleno(C_CABECERA)
    for termino, significado in terminos:
        r += 1
        a = ws.cell(row=r, column=1, value=termino)
        a.font = _f(10, True)
        b = ws.cell(row=r, column=2, value=significado)
        b.font = _f(10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.alignment = Alignment(vertical="top")
        for c in (a, b):
            c.border = BORDE_CELDA
    r += 2
    ws.cell(row=r, column=1, value="Parámetros del análisis").font = _f(12, True)
    r += 1
    for j, titulo in enumerate(("Parámetro", "Valor"), 1):
        c = ws.cell(row=r, column=j, value=titulo)
        c.font = _f(10, True, "FFFFFF")
        c.fill = _relleno(C_CABECERA)
    for k, v in parametros:
        r += 1
        a = ws.cell(row=r, column=1, value=str(k))
        b = ws.cell(row=r, column=2, value=v if isinstance(v, (int, float)) else str(v))
        a.font, b.font = _f(10, True), _f(10)
        b.alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
        a.alignment = Alignment(vertical="top", wrap_text=True)
        for c in (a, b):
            c.border = BORDE_CELDA


# ------------------------------------------------------------------ graficos
def _graficos(wb, datos, ocup):
    """Escribe los datos de los graficos en una hoja oculta y devuelve los graficos."""
    g = wb.create_sheet("datos_graficos")
    g.sheet_state = "hidden"
    graficos = []
    col = 1

    def bloque(titulos, filas):
        nonlocal col
        c0 = col
        for j, t in enumerate(titulos):
            g.cell(row=1, column=c0 + j, value=t)
        for i, fila in enumerate(filas, 2):
            for j, v in enumerate(fila):
                g.cell(row=i, column=c0 + j, value=v)
        col += len(titulos) + 1
        return c0, 1, 1 + len(filas)

    camiones = [c for c in datos.camiones if c.get("tiempo_visible_s")]
    # 1) Movimiento vs detenido por camion
    if camiones:
        c0, r0, r1 = bloque(["Camión", "En movimiento", "Detenido"],
                            [[f"Camión {c['id_camion']}", round((c.get('tiempo_en_movimiento_s') or 0) / 60, 2),
                              round((c.get('tiempo_detenido_s') or 0) / 60, 2)] for c in camiones])
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 55
        ch.add_data(Reference(g, min_col=c0 + 1, max_col=c0 + 2, min_row=r0, max_row=r1), titles_from_data=True)
        ch.set_categories(Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1))
        _color_serie(ch.series[0], C_MOVIMIENTO)
        _color_serie(ch.series[1], C_DETENIDO)
        _titulo(ch, "Tiempo en movimiento y detenido por camión")
        _ejes_limpios(ch)
        _titulo_eje(ch.y_axis, "minutos")
        graficos.append(ch)

    # 2) Reparto del tiempo de la flota
    reparto = [(n, round(s / 60, 2)) for n, s in datos.reparto_tiempo() if s > 0]
    if reparto:
        c0, r0, r1 = bloque(["Estado", "Minutos"], [list(x) for x in reparto])
        ch = DoughnutChart(holeSize=58)
        ch.add_data(Reference(g, min_col=c0 + 1, min_row=r0, max_row=r1), titles_from_data=True)
        ch.set_categories(Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1))
        colores = {"En movimiento": C_MOVIMIENTO, "Detenido en una zona": C_DETENIDO,
                   "Detenido fuera de zonas": C_DETENIDO_FUERA}
        serie = ch.series[0]
        for i, (nombre, _) in enumerate(reparto):
            p = DataPoint(idx=i)
            p.graphicalProperties.solidFill = colores[nombre]
            p.graphicalProperties.line.solidFill = "FFFFFF"
            p.graphicalProperties.line.width = 19050
            serie.dPt.append(p)
        serie.dLbls = DataLabelList()
        serie.dLbls.showPercent = True
        for atributo in ("showVal", "showCatName", "showSerName", "showLeaderLines", "showLegendKey"):
            setattr(serie.dLbls, atributo, False)
        serie.dLbls.txPr = _texto_grafico("", 10, True, "FFFFFF")
        serie.dLbls.txPr.p[0].r = []
        _titulo(ch, "¿En qué se fue el tiempo de la flota?")
        ch.legend.position = "r"
        ch.legend.txPr = _estilo_texto_ejes(9)
        graficos.append(ch)

    # 3) Distancia recorrida por camion
    if camiones:
        clave = "distancia_m" if datos.metros else "distancia_px"
        c0, r0, r1 = bloque(["Camión", f"Distancia ({datos.unidad_d})"],
                            [[f"Camión {c['id_camion']}", round(c.get(clave) or 0, 1)] for c in camiones])
        ch = BarChart()
        ch.type, ch.gapWidth = "col", 60
        ch.add_data(Reference(g, min_col=c0 + 1, min_row=r0, max_row=r1), titles_from_data=True)
        ch.set_categories(Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1))
        _color_serie(ch.series[0], C_MOVIMIENTO)
        ch.legend = None
        _titulo(ch, "Distancia recorrida por camión")
        _ejes_limpios(ch)
        _titulo_eje(ch.y_axis, "metros" if datos.metros else "píxeles")
        ch.series[0].dLbls = DataLabelList()
        ch.series[0].dLbls.showVal = True
        for atributo in ("showCatName", "showSerName", "showPercent", "showLegendKey"):
            setattr(ch.series[0].dLbls, atributo, False)
        ch.series[0].dLbls.numFmt = "#,##0"
        ch.series[0].dLbls.position = "outEnd"
        ch.series[0].dLbls.txPr = _estilo_texto_ejes(8)
        graficos.append(ch)

    # 4) Avisos por tipo
    conteo = Counter(TIPOS_AVISO.get(a.get("tipo"), a.get("titulo", "Otro")) for a in datos.alertas)
    if conteo:
        filas = sorted(conteo.items(), key=lambda x: x[1])
        c0, r0, r1 = bloque(["Tipo", "Avisos"], [list(x) for x in filas])
        ch = BarChart()
        ch.type, ch.gapWidth = "bar", 50
        ch.add_data(Reference(g, min_col=c0 + 1, min_row=r0, max_row=r1), titles_from_data=True)
        ch.set_categories(Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1))
        _color_serie(ch.series[0], SERIES[2])
        ch.legend = None
        _titulo(ch, "Avisos durante el video, por tipo")
        _ejes_limpios(ch)
        ch.x_axis.majorGridlines = None
        ch.series[0].dLbls = DataLabelList()
        ch.series[0].dLbls.showVal = True
        for atributo in ("showCatName", "showSerName", "showPercent", "showLegendKey"):
            setattr(ch.series[0].dLbls, atributo, False)
        ch.series[0].dLbls.position = "outEnd"
        ch.series[0].dLbls.txPr = _estilo_texto_ejes(8)
        graficos.append(ch)

    # 5) Ocupacion de zonas en el tiempo (promedio por intervalo)
    paso, zonas, filas = ocup
    if filas and zonas:
        titulos = ["Desde"] + [datos.nombre_zona(int(z[1:]) - 1) for z in zonas]
        c0, r0, r1 = bloque(titulos, [[_hms(f[0])] + [round(f[1 + 2 * k], 2) for k in range(len(zonas))]
                                      for f in filas])
        ch = LineChart()
        ch.add_data(Reference(g, min_col=c0 + 1, max_col=c0 + len(zonas), min_row=r0, max_row=r1),
                    titles_from_data=True)
        ch.set_categories(Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1))
        for k, s in enumerate(ch.series):
            _color_serie(s, SERIES[k % len(SERIES)], linea=True)
            s.smooth = False
            s.marker.symbol = "none"
        _titulo(ch, "Camiones dentro de cada zona a lo largo del video")
        _ejes_limpios(ch)
        _titulo_eje(ch.y_axis, "camiones (promedio)")
        _titulo_eje(ch.x_axis, f"tiempo del video (intervalos de {paso} s)")
        ch.x_axis.tickLblSkip = max(1, len(filas) // 8)
        graficos.append(ch)

    # 6) Mapa de recorridos
    if datos.tray:
        por = defaultdict(list)
        for f in datos.tray:
            if f.get(datos.eje_x) is not None:
                por[f["id_camion"]].append((f["tiempo_s"], f[datos.eje_x], f[datos.eje_y]))
        ch = ScatterChart()
        ch.scatterStyle = "lineMarker"
        todos_x, todos_y = [], []
        for k, (pid, pts) in enumerate(sorted(por.items())):
            paso_m = max(1, len(pts) // 400)
            pts = pts[::paso_m]
            invertir = not datos.metros or datos.eje_x == "centro_x"   # pixeles: y hacia abajo
            filas_pts, anterior = [], None
            for t, x, y in pts:
                # Si no se vio un rato (salio, estuvo tapado), la linea se corta: no se inventa el camino.
                if anterior is not None and t - anterior > 2.0 + paso_m * datos.intervalo:
                    filas_pts.append([None, None])
                filas_pts.append([x, -y if invertir else y])
                anterior = t
            pts = [p for p in filas_pts if p[0] is not None]
            c0, r0, r1 = bloque([f"X camión {pid}", f"Camión {pid}"], filas_pts)
            s = Series(Reference(g, min_col=c0 + 1, min_row=r0 + 1, max_row=r1),
                       Reference(g, min_col=c0, min_row=r0 + 1, max_row=r1), title=f"Camión {pid}")
            s.marker.symbol = "none"
            s.smooth = False
            _color_serie(s, SERIES[k % len(SERIES)], linea=True, ancho=19050)
            ch.series.append(s)
            todos_x += [p[0] for p in pts]
            todos_y += [p[1] for p in pts]
        ch.display_blanks = "gap"
        _titulo(ch, "Mapa de recorridos (vista desde arriba)")
        _ejes_limpios(ch, rejilla_y=True)
        ch.x_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="EEF0F2")))
        ex, ey = datos.etiquetas_xy
        _titulo_eje(ch.x_axis, ex)
        _titulo_eje(ch.y_axis, ey)
        ancho, alto = datos.d.get("resolucion") or (0, 0)
        if datos.eje_x == "x_m" and datos.georef is not None and ancho:
            esc = datos.georef.escala_m_px
            ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, round(ancho * esc)
            ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, round(alto * esc)
        elif todos_x:
            ch.x_axis.scaling.min, ch.x_axis.scaling.max = math.floor(min(todos_x)), math.ceil(max(todos_x))
            ch.y_axis.scaling.min, ch.y_axis.scaling.max = math.floor(min(todos_y)), math.ceil(max(todos_y))
        ch.x_axis.tickLblPos = "low"
        ch.y_axis.tickLblPos = "low"
        ch.y_axis.number_format = "0;0"            # en pixeles la Y va invertida: sin signo menos
        ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="EEF0F2")))
        ch.legend.position = "r"
        graficos.append(ch)
    return graficos


# ------------------------------------------------------------------ entrada
def generar(datos_registro, ruta, parametros=None, aviso_posiciones=None):
    """Crea el reporte en 'ruta'. datos_registro: dict con video, duracion_s, georef, camiones,
    zonas, eventos, alertas, trayectorias, ocupacion, nombres_zonas, tipos_zonas, intervalo,
    resolucion, limite_general."""
    datos = _Datos(datos_registro)
    wb = Workbook()
    wb.properties.title = TITULO
    wb.properties.creator = ENCABEZADO
    # Primero las hojas de datos (el resumen usa sus referencias); luego se ordenan.
    resumen = wb.active
    ref = _hoja_camiones(wb, datos)
    _hoja_detenciones(wb, datos)
    _hoja_zonas(wb, datos)
    ocup = _hoja_ocupacion(wb, datos)
    _hoja_bitacora(wb, datos)
    _hoja_posiciones(wb, datos, aviso_posiciones)
    _hoja_glosario(wb, datos, parametros or [])
    graficos = _graficos(wb, datos, ocup)
    wb.active = 0
    _hoja_resumen(wb, datos, ref, graficos)
    resumen.sheet_properties.tabColor = C_ACENTO
    for ws in wb.worksheets:
        if ws.title != "Resumen" and ws.sheet_state == "visible":
            ws.sheet_properties.tabColor = C_CABECERA
    wb.save(ruta)
    return ruta
